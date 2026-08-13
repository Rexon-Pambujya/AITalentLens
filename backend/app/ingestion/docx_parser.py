"""DOCX text extraction (section 7: python-docx)."""
import io

import docx

from app.core.exceptions import UnsupportedDocumentError


def extract_text_from_docx(content: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise UnsupportedDocumentError("Could not open DOCX - file may be corrupt") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    # Resumes frequently use tables for layout (e.g. skills grids,
    # contact-info blocks) - skipping tables would silently drop content.
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    text = "\n".join(parts).strip()
    if not text:
        raise UnsupportedDocumentError("No extractable text found in DOCX")
    return text
